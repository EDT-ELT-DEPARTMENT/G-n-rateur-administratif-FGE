# -*- coding: utf-8 -*-
"""Génération d'attestations administratives FR/AR pour la Faculté de Génie Électrique."""
from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from pathlib import Path
import re
import unicodedata
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE_DIR = Path(__file__).resolve().parent

def _resolve_logo_path() -> Path:
    """Priorité au logo.PNG placé à la racine du dépôt GitHub.

    Des chemins de repli sont conservés pour les anciennes installations.
    La casse du nom est volontairement exacte pour Streamlit Cloud sous Linux.
    """
    candidates = (
        BASE_DIR / "logo.PNG",
        BASE_DIR / "logo.png",
        BASE_DIR / "assets" / "logo.PNG",
        BASE_DIR / "assets" / "logo.png",
        BASE_DIR / "assets" / "logo_udl.png",
    )
    return next((path for path in candidates if path.is_file()), candidates[-1])

LOGO_PATH = _resolve_logo_path()

DEPARTEMENTS = {
    "Electrotechnique": {"ar": "الكهروتقني", "fr": "Électrotechnique"},
    "Automatique": {"ar": "الأوتوماتيك", "fr": "Automatique"},
    "Electronique": {"ar": "الإلكترونيك", "fr": "Électronique"},
    "Télécommunications": {"ar": "الاتصالات السلكية واللاسلكية", "fr": "Télécommunications"},
}

FIELD_DEFS = [
    ("department", "Département", ["departement", "département", "department", "dept", "filiere", "section departement"]),
    ("nom", "Nom / Nom de famille", ["nom", "nom de famille", "lastname", "last name", "family name", "surname"]),
    ("prenom", "Prénom", ["prenom", "prénom", "first name", "firstname", "given name"]),
    ("date_naissance", "Date de naissance", ["date naissance", "date de naissance", "date_naissance", "date naissance", "birth date", "date of birth", "dob"]),
    ("lieu_naissance", "Lieu de naissance", ["lieu naissance", "lieu de naissance", "lieu_naissance", "birth place", "place of birth"]),
    ("chef_fr", "Nom du chef de département (français)", ["chef departement fr", "chef de departement", "chef du departement", "responsable departement", "nom chef departement", "chef_fr", "head of department"]),
    ("chef_ar", "Nom du chef de département (arabe)", ["chef departement ar", "chef_ar", "اسم رئيس القسم", "رئيس القسم"]),
    ("numero", "Numéro de l'attestation", ["numero", "numéro", "numero attestation", "numéro attestation", "n attestation", "no attestation"]),
    ("date_attestation", "Date de délivrance", ["date attestation", "date de delivrance", "date_delivrance", "date_attestation", "date emission"]),
    ("reference", "Référence du document", ["reference", "référence", "ref", "reference document", "ref document"]),
    ("matiere_s1", "Matière(s) du semestre 1", ["matiere s1", "matières s1", "matiere semestre 1", "matières semestre 1", "matiere_s1", "matiere premier semestre", "matière premier semestre", "s1 matiere", "s1 matières"]),
    ("debut_s1", "Début de la période au semestre 1", ["debut s1", "date debut s1", "debut periode s1", "date debut semestre 1", "debut_s1"]),
    ("fin_s1", "Fin de la période au semestre 1", ["fin s1", "date fin s1", "fin periode s1", "date fin semestre 1", "fin_s1"]),
    ("section_s1", "Section / niveau / groupe du semestre 1", ["section s1", "groupe s1", "niveau s1", "section_s1", "departement pedagogique s1"]),
    ("annee_s1", "Année universitaire du semestre 1", ["annee s1", "annee universitaire s1", "annee_s1", "année universitaire s1", "annee universitaire"]),
    ("matiere_s2", "Matière(s) du semestre 2", ["matiere s2", "matières s2", "matiere semestre 2", "matières semestre 2", "matiere_s2", "matiere deuxieme semestre", "s2 matiere", "s2 matières"]),
    ("debut_s2", "Début de la période au semestre 2", ["debut s2", "date debut s2", "debut periode s2", "date debut semestre 2", "debut_s2"]),
    ("fin_s2", "Fin de la période au semestre 2", ["fin s2", "date fin s2", "fin periode s2", "date fin semestre 2", "fin_s2"]),
    ("section_s2", "Section / niveau / groupe du semestre 2", ["section s2", "groupe s2", "niveau s2", "section_s2"]),
    ("annee_s2", "Année universitaire du semestre 2", ["annee s2", "annee universitaire s2", "annee_s2", "année universitaire s2"]),
]


def clean(value: Any) -> str:
    if value is None:
        return ""
    try:
        # pandas.NA and NaN safely stringify poorly; avoid the common null values.
        if value != value:  # NaN
            return ""
    except Exception:
        pass
    if isinstance(value, (date, datetime)):
        return value.strftime("%d/%m/%Y")
    s = str(value).strip()
    if s.lower() in {"nan", "nat", "none", "<na>"}:
        return ""
    # Normalize date strings from Excel's YYYY-MM-DD form when possible.
    m = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})(?: 00:00:00)?", s)
    if m:
        return f"{m.group(3).zfill(2)}/{m.group(2).zfill(2)}/{m.group(1)}"
    return s


def normalize(value: Any) -> str:
    s = clean(value).casefold()
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    # Preserve Arabic letters as well as Latin letters and digits for bilingual source headers.
    return " ".join("".join(ch if ch.isalnum() else " " for ch in s).split())


def guess_column(columns: list[Any], aliases: list[str]) -> str | None:
    normalized = {str(c): normalize(c) for c in columns}
    wanted = [normalize(a) for a in aliases]
    # First prefer exact aliases, then acceptable substring matches.
    for target in wanted:
        for col, value in normalized.items():
            if value == target:
                return col
    for target in wanted:
        if len(target) >= 5:
            for col, value in normalized.items():
                if target in value or value in target:
                    return col
    return None


def set_paragraph_direction(paragraph, rtl: bool = False):
    pPr = paragraph._p.get_or_add_pPr()
    bidi = pPr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        pPr.append(bidi)
    bidi.set(qn("w:val"), "1" if rtl else "0")


def set_run_font(run, font: str = "Arial", size: float = 11, bold: bool | None = None, rtl: bool = False):
    run.font.name = font
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.rFonts
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
        rFonts.set(qn(f"w:{attr}"), font)
    if rtl:
        rtl_el = rPr.find(qn("w:rtl"))
        if rtl_el is None:
            rtl_el = OxmlElement("w:rtl")
            rPr.append(rtl_el)
        rtl_el.set(qn("w:val"), "1")
        lang = rPr.find(qn("w:lang"))
        if lang is None:
            lang = OxmlElement("w:lang")
            rPr.append(lang)
        lang.set(qn("w:cs"), "ar-DZ")


def style_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.RIGHT, rtl=False, before=0, after=3, line_spacing=1.0):
    # In Word/LibreOffice, w:bidi reverses the physical meaning of left/right justification.
    # Invert it here so RTL paragraphs are visually positioned at the intended edge.
    if rtl and align == WD_ALIGN_PARAGRAPH.RIGHT:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    elif rtl and align == WD_ALIGN_PARAGRAPH.LEFT:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line_spacing
    set_paragraph_direction(paragraph, rtl=rtl)


def add_para(doc, text="", *, align=WD_ALIGN_PARAGRAPH.RIGHT, rtl=False, size=11, bold=False,
             before=0, after=3, line_spacing=1.0, font="Arial", keep=False):
    p = doc.add_paragraph()
    style_paragraph(p, align, rtl, before, after, line_spacing)
    if text:
        r = p.add_run(str(text))
        set_run_font(r, font=font, size=size, bold=bold, rtl=rtl)
    if keep:
        p.paragraph_format.keep_together = True
    return p


def set_table_borders_none(table):
    tblPr = table._tbl.tblPr
    borders = tblPr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "nil")


def set_cell_padding(cell, top=0, start=50, bottom=0, end=50):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_rtl(table, rtl=False):
    tblPr = table._tbl.tblPr
    existing = tblPr.find(qn("w:bidiVisual"))
    if existing is None:
        existing = OxmlElement("w:bidiVisual")
        tblPr.append(existing)
    existing.set(qn("w:val"), "1" if rtl else "0")


def clear_paragraph(paragraph):
    paragraph.clear()


def setup_section(section):
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(0.9)
    section.bottom_margin = Cm(1.15)
    section.left_margin = Cm(1.25)
    section.right_margin = Cm(1.25)
    section.header_distance = Cm(0.45)
    section.footer_distance = Cm(0.45)


def add_institution_header(doc, language: str, department: str):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(3.05)
    table.columns[1].width = Cm(15.2)
    set_table_borders_none(table)
    set_table_rtl(table, False)
    logo_cell, txt_cell = table.rows[0].cells
    logo_cell.width = Cm(3.05)
    txt_cell.width = Cm(15.2)
    logo_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    txt_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_padding(logo_cell, top=0, start=0, bottom=0, end=100)
    set_cell_padding(txt_cell, top=0, start=0, bottom=0, end=0)
    p0 = logo_cell.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p0.paragraph_format.space_after = Pt(0)
    if LOGO_PATH.exists():
        p0.add_run().add_picture(str(LOGO_PATH), width=Cm(2.9))
    p = txt_cell.paragraphs[0]
    style_paragraph(p, WD_ALIGN_PARAGRAPH.CENTER, rtl=(language == "ar"), after=1, line_spacing=1.0)
    if language == "ar":
        lines = [
            "الجمهورية الجزائرية الديمقراطية الشعبية",
            "وزارة التعليم العالي والبحث العلمي",
            "جامعة جيلالي ليابس - سيدي بلعباس",
            "كلية الهندسة الكهربائية",
        ]
    else:
        lines = [
            "RÉPUBLIQUE ALGÉRIENNE DÉMOCRATIQUE ET POPULAIRE",
            "MINISTÈRE DE L’ENSEIGNEMENT SUPÉRIEUR ET DE LA RECHERCHE SCIENTIFIQUE",
            "UNIVERSITÉ DJILLALI LIABÈS - SIDI BEL ABBÈS",
            "FACULTÉ DE GÉNIE ÉLECTRIQUE",
        ]
    for i, line in enumerate(lines):
        if i:
            p.add_run().add_break()
        r = p.add_run(line)
        set_run_font(r, size=10 if i == 1 else 11, bold=True, rtl=(language == "ar"))
    add_para(doc, "", align=WD_ALIGN_PARAGRAPH.RIGHT, rtl=(language == "ar"), size=3, after=1)
    if language == "ar":
        add_para(doc, f"قسم {DEPARTEMENTS[department]['ar']}", rtl=True, size=11, bold=True, after=1)
    else:
        add_para(doc, f"Département : {DEPARTEMENTS[department]['fr']}", rtl=False, align=WD_ALIGN_PARAGRAPH.RIGHT, size=10.5, bold=True, after=1)


def add_footer(section, reference: str, page_label: str):
    footer = section.footer
    # Unlink the footer from previous section so page labels stay distinct in bilingual output.
    footer.is_linked_to_previous = False
    # Clear default paragraph then create a two-cell reference/footer line.
    if footer.paragraphs:
        footer.paragraphs[0].text = ""
    table = footer.add_table(rows=1, cols=2, width=Cm(18.4))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(15.8)
    table.columns[1].width = Cm(2.6)
    set_table_borders_none(table)
    left, right = table.rows[0].cells
    left.width = Cm(15.8)
    right.width = Cm(2.6)
    p = left.paragraphs[0]
    style_paragraph(p, WD_ALIGN_PARAGRAPH.CENTER, rtl=False, after=0)
    r = p.add_run(f"Réf : {reference}")
    set_run_font(r, size=9)
    p2 = right.paragraphs[0]
    style_paragraph(p2, WD_ALIGN_PARAGRAPH.RIGHT, rtl=False, after=0)
    r2 = p2.add_run(page_label)
    set_run_font(r2, size=9)


def format_value(value: Any, placeholder: str = "................................") -> str:
    return clean(value) or placeholder


def subject_lines(value: Any, max_lines: int = 3) -> list[str]:
    text = clean(value)
    if not text:
        return []
    parts = [p.strip(" \t•-;") for p in re.split(r"[\n;|]+", text) if p.strip(" \t•-;")]
    # Keep every subject even though the scanned form provides only three dotted lines.
    if len(parts) > max_lines:
        parts = parts[:max_lines - 1] + [" ; ".join(parts[max_lines - 1:])]
    return parts


def issue_year(values: dict[str, Any]) -> str:
    issue_date = clean(values.get("date_attestation"))
    match = re.search(r"(?:19|20)\d{2}", issue_date)
    return match.group(0) if match else str(date.today().year)


def add_field_pair(doc, left_label: str, left_value: str, right_label: str, *, size=10.5, value_rtl: bool = False):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Cm(8.8)
    table.columns[1].width = Cm(9.0)
    set_table_borders_none(table)
    left, right = table.rows[0].cells
    left.width = Cm(8.8)
    right.width = Cm(9.0)
    for c in (left, right):
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_padding(c, top=0, start=0, bottom=0, end=100)
    p = left.paragraphs[0]
    style_paragraph(p, WD_ALIGN_PARAGRAPH.LEFT, rtl=False, after=1)
    r = p.add_run(f"{left_label} : ")
    set_run_font(r, size=size)
    value_run = p.add_run(left_value)
    set_run_font(value_run, size=size, rtl=value_rtl)
    p = right.paragraphs[0]
    style_paragraph(p, WD_ALIGN_PARAGRAPH.RIGHT, rtl=True, after=1)
    r = p.add_run(right_label)
    set_run_font(r, size=size, rtl=True)
    return table


def add_semester_ar(doc, semester_num: int, values: dict[str, Any]):
    key = f"s{semester_num}"
    heading_before = 25 if semester_num == 1 else 15
    add_para(doc, f"السداسي {'الأول' if semester_num == 1 else 'الثاني'} في مادة:", rtl=True, size=12, bold=True, before=heading_before, after=3)
    lines = subject_lines(values.get(f"matiere_{key}"), 3)
    for i in range(3):
        txt = lines[i] if i < len(lines) else "................................................................................................................"
        # The supplied paper model leaves generous space for writing subject names.
        add_para(doc, txt, rtl=True, size=10.5, after=23, line_spacing=1.0)
    period = f"تاريخ التوظيف: ابتداءً من {format_value(values.get(f'debut_{key}'), '……/……/………')} إلى غاية {format_value(values.get(f'fin_{key}'), '……/……/………')}"
    add_para(doc, period, rtl=True, size=10.5, before=2, after=8)
    section = format_value(values.get(f"section_{key}"), "................")
    year = format_value(values.get(f"annee_{key}"), "……/……")
    add_para(doc, f"القسم: {section}       السنة الجامعية: {year}", rtl=True, size=10.5, after=15)


def add_arabic_page(doc, values: dict[str, Any], page_label: str):
    section = doc.sections[-1]
    setup_section(section)
    dept = values.get("department", "Electrotechnique")
    add_institution_header(doc, "ar", dept)
    # Department and issue details mirror the upper-right block of the scanned form.
    add_para(doc, f"رقم: {format_value(values.get('numero'), '……')} / {issue_year(values)}", rtl=True, size=10.5, after=0)
    p = add_para(doc, "شهادة إدارية", align=WD_ALIGN_PARAGRAPH.CENTER, rtl=True, size=18, bold=True, before=40, after=30)
    p.paragraph_format.keep_with_next = True
    chef = format_value(values.get("chef_ar") or values.get("chef_fr"), "................................")
    dept_ar = DEPARTEMENTS[dept]["ar"]
    sentence = f"أنا الممضي(ة) أسفله السيد(ة) {chef} رئيس(ة) قسم {dept_ar} بكلية الهندسة الكهربائية بجامعة جيلالي ليابس بسيدي بلعباس، أشهد أن السيد(ة):"
    add_para(doc, sentence, rtl=True, size=11, after=12, line_spacing=1.0)
    nom_ar_value = format_value(values.get("nom_ar") or values.get("nom"))
    prenom_ar_value = format_value(values.get("prenom_ar") or values.get("prenom"))
    add_field_pair(doc, "Nom", nom_ar_value, "اللقب:", value_rtl=bool(values.get("nom_ar")))
    add_field_pair(doc, "Prénom", prenom_ar_value, "الاسم:", value_rtl=bool(values.get("prenom_ar")))
    birth_date = format_value(values.get("date_naissance"), "……/……/………")
    birth_place = format_value(values.get("lieu_naissance"), "........................")
    add_para(doc, f"المولود(ة) بتاريخ: {birth_date} بـ {birth_place}", rtl=True, size=10.5, after=50)
    add_semester_ar(doc, 1, values)
    add_semester_ar(doc, 2, values)
    add_para(doc, "سُلّمت هذه الشهادة بطلب من المعني(ة) بالأمر لاستعمالها في حدود ما يسمح به القانون.", rtl=True, size=10.5, bold=True, before=8, after=4)
    add_footer(section, values.get("reference", "UDL-GEL-ER-120-2026"), page_label)


def add_semester_fr(doc, n: int, values: dict[str, Any]):
    key = f"s{n}"
    label = "Premier semestre" if n == 1 else "Deuxième semestre"
    before = 5 if n == 1 else 25
    add_para(doc, f"{label} — Matière(s) enseignée(s) :", align=WD_ALIGN_PARAGRAPH.LEFT, size=11.5, bold=True, before=before, after=3)
    lines = subject_lines(values.get(f"matiere_{key}"), 3)
    for i in range(3):
        txt = lines[i] if i < len(lines) else "........................................................................................................"
        add_para(doc, txt, align=WD_ALIGN_PARAGRAPH.LEFT, size=10.5, after=8)
    add_para(doc, f"Période d'enseignement : du {format_value(values.get(f'debut_{key}'), '……/……/………')} au {format_value(values.get(f'fin_{key}'), '……/……/………')}", align=WD_ALIGN_PARAGRAPH.LEFT, size=10.5, after=6)
    add_para(doc, f"Groupe / section : {format_value(values.get(f'section_{key}'), '................')}       Année universitaire : {format_value(values.get(f'annee_{key}'), '……/……')}", align=WD_ALIGN_PARAGRAPH.LEFT, size=10.5, after=8)


def add_french_page(doc, values: dict[str, Any], page_label: str):
    section = doc.sections[-1]
    setup_section(section)
    dept = values.get("department", "Electrotechnique")
    add_institution_header(doc, "fr", dept)
    add_para(doc, f"N° : {format_value(values.get('numero'), '……')} / {issue_year(values)}", align=WD_ALIGN_PARAGRAPH.RIGHT, size=10.5, after=0)
    add_para(doc, "ATTESTATION ADMINISTRATIVE", align=WD_ALIGN_PARAGRAPH.CENTER, size=17, bold=True, before=65, after=16)
    chef = format_value(values.get("chef_fr") or values.get("chef_ar"), "................................")
    dept_fr = DEPARTEMENTS[dept]["fr"]
    body = (f"Je soussigné(e), {chef}, Chef du département {dept_fr}, "
            "à la Faculté de Génie Électrique de l'Université Djillali Liabès de Sidi Bel Abbès, "
            "atteste que :")
    add_para(doc, body, align=WD_ALIGN_PARAGRAPH.LEFT, size=11, after=14, line_spacing=1.05)
    add_para(doc, f"Nom : {format_value(values.get('nom'))}", align=WD_ALIGN_PARAGRAPH.LEFT, size=11, after=3)
    add_para(doc, f"Prénom : {format_value(values.get('prenom'))}", align=WD_ALIGN_PARAGRAPH.LEFT, size=11, after=3)
    add_para(doc, f"Né(e) le : {format_value(values.get('date_naissance'), '……/……/………')} à : {format_value(values.get('lieu_naissance'), '........................')}", align=WD_ALIGN_PARAGRAPH.LEFT, size=11, after=6)
    add_semester_fr(doc, 1, values)
    add_semester_fr(doc, 2, values)
    add_para(doc, "La présente attestation est délivrée à l'intéressé(e) à sa demande, pour servir et valoir ce que de droit.", align=WD_ALIGN_PARAGRAPH.LEFT, size=10.5, bold=True, before=5, after=6)
    add_footer(section, values.get("reference", "UDL-GEL-ER-120-2026"), page_label)


def generate_attestation(values: dict[str, Any], language_mode: str = "Bilingue (arabe puis français)") -> bytes:
    doc = Document()
    doc.core_properties.title = "Attestation administrative - Faculté de Génie Électrique"
    doc.core_properties.subject = "Attestation administrative générée à partir du modèle fourni"
    doc.core_properties.author = "Université Djillali Liabès - Sidi Bel Abbès"
    if language_mode.startswith("Arabe"):
        add_arabic_page(doc, values, "1/1")
    elif language_mode.startswith("Français"):
        add_french_page(doc, values, "1/1")
    else:
        add_arabic_page(doc, values, "1/2")
        sec = doc.add_section(WD_SECTION.NEW_PAGE)
        setup_section(sec)
        add_french_page(doc, values, "2/2")
    buff = BytesIO()
    doc.save(buff)
    return buff.getvalue()
