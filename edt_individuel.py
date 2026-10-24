# -*- coding: utf-8 -*-
"""Import, rapprochement et exports d'un emploi du temps individuel."""
from __future__ import annotations

from difflib import SequenceMatcher
from html import escape
from io import BytesIO
import re
import unicodedata
from xml.sax.saxutils import escape as xml_escape

import pandas as pd

DAYS_ORDER = ["Dimanche", "Lundi", "Mardi", "Mercredi", "Jeudi", "Vendredi", "Samedi"]
TYPE_COLORS = {
    "Cours": {"hex": "1D4ED8", "fill": "DBEAFE", "font": "1E40AF"},
    "TD": {"hex": "15803D", "fill": "DCFCE7", "font": "166534"},
    "TP": {"hex": "B91C1C", "fill": "FEE2E2", "font": "991B1B"},
    "Autre": {"hex": "64748B", "fill": "F1F5F9", "font": "334155"},
}


def _strip_accents(value: object) -> str:
    text = "" if value is None else str(value).strip()
    text = unicodedata.normalize("NFKD", text)
    return "".join(ch for ch in text if not unicodedata.combining(ch))


def normalize_header(value: object) -> str:
    return " ".join(re.sub(r"[^a-z0-9]+", " ", _strip_accents(value).casefold()).split())


def normalize_teacher(value: object) -> str:
    text = normalize_header(value)
    # Retire les titres professionnels fréquents présents dans certaines cellules EDT.
    tokens = [t for t in text.split() if t not in {"mr", "mme", "m", "dr", "dre", "pr", "prof", "professeur", "professeure"}]
    return " ".join(tokens)


def _normalize_time_label(value: object) -> str:
    """Unifie les variantes d'écriture du même créneau (11h-12h30 / 11h - 12h30)."""
    text = _strip_accents(value).lower().replace(":", "h")
    parts = re.findall(r"(\d{1,2})\s*h\s*(\d{0,2})", text)
    if len(parts) >= 2:
        normalized = []
        for hour, minute in parts[:2]:
            h = str(int(hour))
            minute_value = int(minute or 0)
            normalized.append(f"{h}h{minute_value:02d}" if minute_value else f"{h}h")
        return " - ".join(normalized)
    return re.sub(r"\s*[-–—]\s*", " - ", str(value or "").strip())


def _time_start(value: object) -> int:
    text = _strip_accents(value).lower().replace(":", "h")
    match = re.search(r"(\d{1,2})\s*h\s*(\d{0,2})", text)
    if not match:
        return 10**8
    return int(match.group(1)) * 60 + int(match.group(2) or 0)


def _day_rank(value: object) -> int:
    key = normalize_header(value)
    for idx, day in enumerate(DAYS_ORDER):
        if key == normalize_header(day):
            return idx
    return 100


def classify_teaching_type(code: object, title: object) -> str:
    code_norm = normalize_header(code)
    title_norm = normalize_header(title)
    haystack = f" {code_norm} {title_norm} "
    # Priorité au code EDT : c'est le champ de classification le plus fiable.
    if re.search(r"(?:^|\s)(?:tp|travaux pratiques?)(?:\s|\d|$)", f"{code_norm} ") or title_norm.startswith("tp ") or title_norm.startswith("tp-") or "travaux pratiques" in haystack:
        return "TP"
    if re.search(r"(?:^|\s)(?:td|travaux diriges?)(?:\s|\d|$)", f"{code_norm} ") or title_norm.startswith("td ") or title_norm.startswith("td-") or "travaux diriges" in haystack:
        return "TD"
    if code_norm.startswith("cours") or title_norm.startswith("cours ") or title_norm.startswith("cours-") or re.search(r"(?:^|\s)cm(?:\s|$)", f" {code_norm} "):
        return "Cours"
    return "Autre"


def extract_subject(title: object, code: object = "") -> str:
    """Nettoie les préfixes de type et les suffixes de groupe sans réécrire la matière."""
    result = "" if title is None else str(title).strip()
    result = re.sub(r"^\s*(?:cours|travaux diriges?|travaux pratiques?|td|tp|cm)\s*[-:–—]\s*", "", result, flags=re.IGNORECASE)
    # Exemples de suffixes de structure présents dans l'EDT : -SG12, -G2, -Section A.
    result = re.sub(r"\s*[-–—]\s*(?:SG\s*\d+|S\s*/?\s*G\s*\d+|GROUPE\s*(?:TD|TP)?\s*[A-Z0-9_-]*|GRP\s*[A-Z0-9_-]*|G\s*\d+|SECTION\s*[A-Z0-9_-]*)\s*$", "", result, flags=re.IGNORECASE)
    return re.sub(r"\s+", " ", result).strip(" -–—")


def split_teacher_cell(value: object) -> list[str]:
    raw = "" if value is None else str(value).strip()
    names = re.split(r"\s*(?:/|\+|;|&|\n)\s*", raw)
    return [re.sub(r"\s+", " ", n).strip(" .;,/") for n in names if n.strip(" .;,/+")]


def _names_match_score(person_name: str, edt_name: str) -> float:
    a = normalize_teacher(person_name).split()
    b = normalize_teacher(edt_name).split()
    if not a or not b:
        return 0.0
    if " ".join(a) == " ".join(b) or sorted(a) == sorted(b):
        return 1.0
    # Appariement prudent : au moins deux éléments nominatifs doivent correspondre,
    # avec tolérance aux accents, ordre différent et petites variantes de translittération.
    remaining = set(range(len(b)))
    matched_a: list[int] = []
    matched_b: list[int] = []
    pair_scores: list[float] = []
    for i, token_a in enumerate(a):
        best = None
        for j in remaining:
            token_b = b[j]
            score = SequenceMatcher(None, token_a, token_b).ratio()
            if token_a == token_b:
                score = 1.0
            if len(token_a) <= 3 or len(token_b) <= 3:
                score = 1.0 if token_a == token_b else 0.0
            if best is None or score > best[0]:
                best = (score, j)
        if best and best[0] >= 0.84:
            matched_a.append(i)
            matched_b.append(best[1])
            pair_scores.append(best[0])
            remaining.remove(best[1])
    if len(matched_a) < 2:
        return 0.0
    coverage_a = len(matched_a) / len(a)
    coverage_b = len(matched_b) / len(b)
    # Une source peut contenir un troisième prénom absent de l'EDT. Si les deux
    # éléments du nom visibles dans l'EDT correspondent de façon certaine, on autorise
    # ce cas (ex. « LOUSDAD Aymen Aziz » vs « LOUSDAD AYMEN »).
    if len(matched_a) >= 2 and coverage_b >= 0.90 and coverage_a >= 0.60:
        return (sum(pair_scores) / len(pair_scores)) * (0.90 + 0.10 * coverage_a)
    if coverage_a < 0.66 or coverage_b < 0.75:
        return 0.0
    return (sum(pair_scores) / len(pair_scores)) * min(coverage_a, coverage_b)


def load_edt_sessions(file_bytes: bytes) -> list[dict[str, str]]:
    if not file_bytes:
        return []
    frame = pd.read_excel(BytesIO(file_bytes), dtype=str).fillna("")
    columns = {normalize_header(c): str(c) for c in frame.columns}
    aliases = {
        "Enseignements": ["enseignements", "enseignement", "matiere", "intitule"],
        "Code": ["code", "code enseignement", "type"],
        "Enseignants": ["enseignants", "enseignant", "professeur", "professeurs"],
        "Horaire": ["horaire", "creneau", "creneaux horaires"],
        "Jours": ["jours", "jour"],
        "Lieu": ["lieu", "salle", "local"],
        "Promotion": ["promotion", "promotions", "groupe", "classe"],
        "Département": ["departement", "departement affiliation", "department"],
    }
    colmap: dict[str, str] = {}
    for key, options in aliases.items():
        for option in options:
            n = normalize_header(option)
            if n in columns:
                colmap[key] = columns[n]
                break
    missing = [col for col in ("Enseignements", "Enseignants", "Horaire", "Jours") if col not in colmap]
    if missing:
        raise ValueError("Colonnes EDT indispensables introuvables : " + ", ".join(missing))
    rows: list[dict[str, str]] = []
    for _, row in frame.iterrows():
        item = {key: str(row.get(source, "") or "").strip() for key, source in colmap.items()}
        if not item.get("Enseignements") or not item.get("Enseignants") or not item.get("Horaire") or not item.get("Jours"):
            continue
        item.setdefault("Code", "")
        item.setdefault("Lieu", "")
        item.setdefault("Promotion", "")
        item.setdefault("Département", "")
        item["Horaire"] = _normalize_time_label(item.get("Horaire", ""))
        item["Type"] = classify_teaching_type(item.get("Code", ""), item["Enseignements"])
        item["Matiere"] = extract_subject(item["Enseignements"], item.get("Code", ""))
        item["Enseignants_liste"] = split_teacher_cell(item["Enseignants"])
        rows.append(item)
    rows.sort(key=lambda x: (_day_rank(x.get("Jours")), _time_start(x.get("Horaire")), normalize_header(x.get("Promotion")), normalize_header(x.get("Matiere"))))
    return rows


def unique_teacher_names(sessions: list[dict]) -> list[str]:
    names: set[str] = set()
    for session in sessions:
        names.update(session.get("Enseignants_liste") or split_teacher_cell(session.get("Enseignants", "")))
    return sorted((n for n in names if n), key=lambda s: normalize_header(s))


def match_teacher_sessions(person_full_name: str, sessions: list[dict], explicit_names: list[str] | None = None) -> tuple[list[dict], list[str], str]:
    """Retourne les séances correspondant au demandeur et les noms EDT retenus."""
    if explicit_names:
        chosen = {normalize_teacher(n) for n in explicit_names}
        found = [s for s in sessions if any(normalize_teacher(n) in chosen for n in (s.get("Enseignants_liste") or []))]
        return found, sorted(explicit_names), "manuel"
    scores: dict[str, float] = {}
    for name in unique_teacher_names(sessions):
        score = _names_match_score(person_full_name, name)
        if score:
            scores[name] = score
    # When ambiguous, keep only strongest score(s); avoids including unrelated names.
    if not scores:
        return [], [], "aucun"
    best_score = max(scores.values())
    if best_score < 0.82:
        return [], [], "aucun"
    if best_score >= 0.99:
        selected_names = sorted([name for name, score in scores.items() if score >= 0.99], key=lambda s: normalize_header(s))
    else:
        selected_names = sorted([name for name, score in scores.items() if score >= max(0.82, best_score - 0.015)], key=lambda s: normalize_header(s))
    chosen = {normalize_teacher(n) for n in selected_names}
    found = [s for s in sessions if any(normalize_teacher(n) in chosen for n in (s.get("Enseignants_liste") or []))]
    return found, selected_names, "exact" if best_score >= 0.99 else "approximatif"


def unique_subjects(sessions: list[dict]) -> list[str]:
    subjects: list[str] = []
    seen: set[str] = set()
    for session in sessions:
        sub = str(session.get("Matiere") or extract_subject(session.get("Enseignements", "")))
        k = normalize_header(sub)
        if sub and k not in seen:
            seen.add(k)
            subjects.append(sub)
    return subjects


def build_grid(sessions: list[dict]) -> tuple[list[str], list[str], dict[tuple[str, str], list[dict]]]:
    days_present = sorted({s.get("Jours", "") for s in sessions if s.get("Jours")}, key=lambda x: (_day_rank(x), normalize_header(x)))
    time_slots = sorted({s.get("Horaire", "") for s in sessions if s.get("Horaire")}, key=lambda x: (_time_start(x), normalize_header(x)))
    grid: dict[tuple[str, str], list[dict]] = {}
    for session in sessions:
        key = (session.get("Jours", ""), session.get("Horaire", ""))
        grid.setdefault(key, []).append(session)
    for key in grid:
        grid[key].sort(key=lambda x: (normalize_header(x.get("Promotion")), normalize_header(x.get("Matiere"))))
    return days_present, time_slots, grid


def _cell_plain(items: list[dict]) -> str:
    lines = []
    for item in items:
        parts = [f"[{item.get('Type', 'Autre')}] {item.get('Matiere') or item.get('Enseignements', '')}"]
        detail = " · ".join(x for x in [item.get("Promotion", ""), item.get("Lieu", "")] if x)
        if detail:
            parts.append(detail)
        lines.append("\n".join(parts))
    return "\n\n---\n\n".join(lines)


def export_timetable_html(teacher_name: str, sessions: list[dict]) -> bytes:
    days, times, grid = build_grid(sessions)
    headers = "".join(f"<th>{escape(t)}</th>" for t in times)
    body_rows = []
    for day in days:
        cells = []
        for time in times:
            items = grid.get((day, time), [])
            if not items:
                cells.append("<td class='empty'></td>")
                continue
            blocks = []
            for item in items:
                typ = item.get("Type", "Autre")
                color = TYPE_COLORS.get(typ, TYPE_COLORS["Autre"])
                detail = " · ".join(x for x in [item.get("Promotion", ""), item.get("Lieu", "")] if x)
                blocks.append(
                    f"<div class='lesson {typ.lower()}' style='--accent:#{color['hex']};--pale:#{color['fill']}'>"
                    f"<b>{escape(typ.upper())}</b><br><strong>{escape(item.get('Matiere') or item.get('Enseignements',''))}</strong>"
                    f"{('<br><span>'+escape(detail)+'</span>') if detail else ''}</div>"
                )
            cells.append("<td>" + "".join(blocks) + "</td>")
        body_rows.append(f"<tr><th class='day'>{escape(day)}</th>{''.join(cells)}</tr>")
    title = escape(f"Emploi du temps individuel — {teacher_name}")
    html = f"""<!doctype html><html lang='fr'><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'><title>{title}</title>
<style>body{{font-family:Arial,sans-serif;margin:20px;background:#f1f5f9;color:#0f172a}}h1{{color:#1e3a8a;font-size:24px}}p{{color:#475569}}.wrap{{overflow:auto;background:white;padding:10px;border-radius:12px}}table{{border-collapse:separate;border-spacing:4px;width:100%;min-width:900px;table-layout:fixed}}th,td{{vertical-align:top;border:1px solid #cbd5e1;border-radius:6px;padding:8px;overflow-wrap:anywhere}}thead th{{background:#1d4ed8;color:white;text-align:center;position:sticky;top:0}}th.day{{background:#0f172a;color:white;width:105px}}td{{background:#f8fafc;min-width:120px}}td.empty{{background:#f8fafc}}.lesson{{border-left:5px solid var(--accent);background:var(--pale);padding:7px;margin-bottom:6px;border-radius:4px;font-size:12px}}.lesson b{{font-size:10px;color:var(--accent)}}.lesson strong{{display:inline-block;margin-top:3px}}.lesson span{{display:inline-block;margin-top:4px;color:#334155}}@media print{{body{{margin:4mm}}.wrap{{overflow:visible}}table{{min-width:0}}}}</style></head><body>
<h1>{title}</h1><p>Jours en lignes verticales · Horaires en colonnes horizontales · Source : EDT_Edition.xlsx · {len(sessions)} séance(s)</p>
<div class='wrap'><table><thead><tr><th>Jour / Horaire</th>{headers}</tr></thead><tbody>{''.join(body_rows)}</tbody></table></div></body></html>"""
    return html.encode("utf-8")


def export_timetable_xlsx(teacher_name: str, sessions: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.page import PageMargins

    days, times, grid = build_grid(sessions)
    wb = Workbook()
    ws = wb.active
    ws.title = "EDT individuel"
    last_col = max(2, len(times) + 1)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    title = ws.cell(1, 1, f"EMPLOI DU TEMPS INDIVIDUEL — {teacher_name}")
    title.font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
    title.fill = PatternFill("solid", fgColor="1E3A8A")
    title.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    ws.cell(2, 1, f"Source : EDT_Edition.xlsx | {len(sessions)} séance(s) | Généré automatiquement").font = Font(italic=True, color="475569", size=9)
    headers = ["Jour / Horaire"] + times
    for col, value in enumerate(headers, start=1):
        cell = ws.cell(4, col, value)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1D4ED8")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="thin", color="FFFFFF"))
    ws.row_dimensions[4].height = 28
    for row_idx, day in enumerate(days, start=5):
        dcell = ws.cell(row_idx, 1, day)
        dcell.font = Font(bold=True, color="FFFFFF")
        dcell.fill = PatternFill("solid", fgColor="0F172A")
        dcell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for col_idx, slot in enumerate(times, start=2):
            items = grid.get((day, slot), [])
            cell = ws.cell(row_idx, col_idx, _cell_plain(items))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(left=Side(style="thin", color="CBD5E1"), right=Side(style="thin", color="CBD5E1"), top=Side(style="thin", color="CBD5E1"), bottom=Side(style="thin", color="CBD5E1"))
            if len({it.get("Type", "Autre") for it in items}) == 1 and items:
                fill = TYPE_COLORS.get(items[0].get("Type", "Autre"), TYPE_COLORS["Autre"])["fill"]
            else:
                fill = "F8FAFC"
            cell.fill = PatternFill("solid", fgColor=fill)
        ws.row_dimensions[row_idx].height = 72 if len(times) <= 8 else 90
    ws.column_dimensions["A"].width = 17
    for i in range(2, last_col + 1):
        ws.column_dimensions[get_column_letter(i)].width = 24 if len(times) <= 7 else 20
    ws.freeze_panes = "B5"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3 if len(times) > 8 else ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.4, bottom=0.4, header=0.15, footer=0.15)
    ws.print_title_rows = "1:4"
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def export_timetable_pdf(teacher_name: str, sessions: list[dict]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    days, times, grid = build_grid(sessions)
    out = BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), rightMargin=8*mm, leftMargin=8*mm, topMargin=8*mm, bottomMargin=8*mm, title=f"EDT individuel - {teacher_name}")
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("TitleCustom", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=15, leading=18, textColor=colors.HexColor("#1E3A8A"), alignment=TA_CENTER, spaceAfter=3*mm)
    meta_style = ParagraphStyle("MetaCustom", parent=styles["Normal"], fontName="Helvetica", fontSize=8, leading=10, textColor=colors.HexColor("#475569"), alignment=TA_CENTER, spaceAfter=4*mm)
    slots_n = max(1, len(times))
    cell_font = 7 if slots_n <= 6 else 6.2 if slots_n <= 8 else 5.5 if slots_n <= 10 else 4.8
    body_style = ParagraphStyle("CellCustom", parent=styles["Normal"], fontName="Helvetica", fontSize=cell_font, leading=cell_font+1.4, alignment=TA_LEFT, wordWrap="CJK")
    header_style = ParagraphStyle("HeaderCustom", parent=body_style, fontName="Helvetica-Bold", textColor=colors.white, alignment=TA_CENTER, fontSize=max(5.2, cell_font))
    day_style = ParagraphStyle("DayCustom", parent=body_style, fontName="Helvetica-Bold", textColor=colors.white, alignment=TA_CENTER)
    story = [Paragraph(xml_escape(f"EMPLOI DU TEMPS INDIVIDUEL — {teacher_name}"), title_style), Paragraph(xml_escape(f"Source : EDT_Edition.xlsx · {len(sessions)} séance(s)"), meta_style)]
    data = [[Paragraph("Jour / Horaire", header_style)] + [Paragraph(xml_escape(t), header_style) for t in times]]
    backgrounds: list[tuple[int,int,str]] = []
    for day in days:
        row: list[object] = [Paragraph(xml_escape(day), day_style)]
        row_index = len(data)
        for col_idx, slot in enumerate(times, start=1):
            items = grid.get((day, slot), [])
            if not items:
                row.append("")
                continue
            lines = []
            types = {i.get("Type", "Autre") for i in items}
            for item in items:
                typ = item.get("Type", "Autre")
                subject = xml_escape(item.get("Matiere") or item.get("Enseignements", ""))
                detail = xml_escape(" · ".join(x for x in [item.get("Promotion", ""), item.get("Lieu", "")] if x))
                lines.append(f"<b>{xml_escape(typ.upper())}</b><br/>{subject}" + (f"<br/><font color='#475569'>{detail}</font>" if detail else ""))
            row.append(Paragraph("<br/><br/>".join(lines), body_style))
            if len(types) == 1:
                c = TYPE_COLORS.get(next(iter(types)), TYPE_COLORS["Autre"])["fill"]
                backgrounds.append((row_index, col_idx, c))
        data.append(row)
    available_width = landscape(A4)[0] - 16*mm
    first_w = 25*mm
    other_w = (available_width - first_w) / max(1, len(times))
    table = Table(data, colWidths=[first_w] + [other_w] * len(times), repeatRows=1, hAlign="LEFT")
    commands = [
        ("BACKGROUND", (0,0), (-1,0), colors.HexColor("#1D4ED8")),
        ("BACKGROUND", (0,1), (0,-1), colors.HexColor("#0F172A")),
        ("GRID", (0,0), (-1,-1), 0.45, colors.HexColor("#CBD5E1")),
        ("VALIGN", (0,0), (-1,-1), "TOP"),
        ("LEFTPADDING", (0,0), (-1,-1), 3), ("RIGHTPADDING", (0,0), (-1,-1), 3),
        ("TOPPADDING", (0,0), (-1,-1), 4), ("BOTTOMPADDING", (0,0), (-1,-1), 4),
        ("ALIGN", (0,0), (-1,0), "CENTER"),
    ]
    for row_i, col_i, fill in backgrounds:
        commands.append(("BACKGROUND", (col_i, row_i), (col_i, row_i), colors.HexColor("#" + fill)))
    table.setStyle(TableStyle(commands))
    story.append(table)
    doc.build(story)
    return out.getvalue()


def timetable_grid_records(sessions: list[dict]) -> tuple[list[str], list[str], list[list[str]]]:
    days, times, grid = build_grid(sessions)
    matrix: list[list[str]] = []
    for day in days:
        matrix.append([day] + [_cell_plain(grid.get((day, slot), [])) for slot in times])
    return days, times, matrix
